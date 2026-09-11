# -*- coding: utf-8 -*-
"""왕복 검증: 생성 xlsx → 더미 채움 → 격리폴더 주입 → 재읽기 대조 → 원상복구."""
import os
import shutil
import subprocess
import sys
import tempfile

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
ERB_DIR = os.path.join(ROOT, "에라마왕 개조판 1.28", "ERB")
XLSX = os.path.join(ROOT, "에라마왕 개조판 1.28", "_작성필요_대사목록.xlsx")

TEST_FILE = "EVENT_K5_マオ.ERB"

tmp = tempfile.mkdtemp(prefix="fill_test_")
tmp_erb = os.path.join(tmp, "erb")
os.makedirs(tmp_erb)
src = os.path.join(ERB_DIR, TEST_FILE)
dst = os.path.join(tmp_erb, TEST_FILE)
shutil.copy2(src, dst)

orig_bytes = open(dst, "rb").read()
had_bom = orig_bytes.startswith(b"\xef\xbb\xbf")
has_crlf = b"\r\n" in orig_bytes
orig_lines = orig_bytes.decode("utf-8-sig").split("\n")

# 테스트용 xlsx 복사 + 대상 파일의 앞쪽 슬롯 3개에 더미 대사 기입
test_xlsx = os.path.join(tmp, "test.xlsx")
shutil.copy2(XLSX, test_xlsx)
wb = load_workbook(test_xlsx)
ws = wb["작성필요"]
header = [c.value for c in ws[1]]
c_file = header.index("파일") + 1
c_input = header.index("작성할 대사 ← 여기에 입력") + 1
c_line = header.index("_line") + 1

targets = []
for r in range(2, ws.max_row + 1):
    if ws.cell(r, c_file).value == TEST_FILE:
        line = int(ws.cell(r, c_line).value)
        marker = f"[[TESTMARK {len(targets)} 마オ테스트]]"
        ws.cell(r, c_input).value = marker
        targets.append((line, marker))
        if len(targets) == 3:
            break
wb.save(test_xlsx)
assert targets, "테스트 대상 슬롯을 찾지 못함"

# 격리폴더로 주입
env = dict(os.environ, FILL_ERB_DIR=tmp_erb)
res = subprocess.run(
    [sys.executable, os.path.join(HERE, "apply_fill_list.py"), test_xlsx],
    capture_output=True, text=True, env=env,
)
print(res.stdout)
if res.returncode != 0:
    print(res.stderr); sys.exit(1)

# 재읽기 대조
new_bytes = open(dst, "rb").read()
assert new_bytes.startswith(b"\xef\xbb\xbf") == had_bom, "BOM 보존 실패"
assert (b"\r\n" in new_bytes) == has_crlf, "CRLF 보존 실패"
new_lines = new_bytes.decode("utf-8-sig").split("\n")
assert len(new_lines) == len(orig_lines), f"라인 수 변경됨 {len(orig_lines)}->{len(new_lines)}"

changed = 0
for line, marker in targets:
    got = new_lines[line - 1].rstrip("\r")
    assert marker in got, f"라인 {line} 주입 실패: {got!r}"
    assert got.lstrip().startswith("PRINTFORM"), f"명령어 손상: {got!r}"
    assert got.startswith("\t") or got.startswith(orig_lines[line - 1][:1]), "들여쓰기 손상"
    changed += 1

# 비대상 라인 무변경 확인
target_lines = {t[0] for t in targets}
for idx in range(len(orig_lines)):
    if (idx + 1) in target_lines:
        continue
    assert orig_lines[idx] == new_lines[idx], f"비대상 라인 {idx+1} 변경됨"

print(f"PASS: {changed}개 주입, BOM={had_bom}, CRLF={has_crlf}, 라인수={len(new_lines)} 유지, 비대상 무변경")
shutil.rmtree(tmp, ignore_errors=True)
