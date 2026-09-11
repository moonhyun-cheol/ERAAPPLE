# -*- coding: utf-8 -*-
"""
빈 PRINTFORM 슬롯(대사 미작성 자리)을 전량 스캔해
'캐릭터별 시트' 엑셀(.xlsx)로 추출한다.

- 시트 1개 = 캐릭터(파일) 1개.
- 각 시트 맨 위에 그 캐릭터가 어떤 인물인지 소개(성격/성향/말투)를 넣어,
  대사를 쓸 때 톤을 잡기 쉽게 한다.
- 노란 '작성할 대사' 칸에만 입력하면 apply_fill_list.py 가 원본 자리에 주입한다.

사용법:
    python tools/fill-list/build_fill_list.py
"""
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
ERB_DIR = os.path.join(ROOT, "에라마왕 개조판 1.28", "ERB")
OUT_XLSX = os.path.join(ROOT, "에라마왕 개조판 1.28", "_작성필요_대사목록.xlsx")

EMPTY_FORM = re.compile(
    r"^(PRINTFORMW|PRINTFORML|PRINTFORMLC|PRINTFORMC|PRINTFORMK|PRINTFORMD|PRINTSINGLEFORM|PRINTFORM)\s*$"
)
DIVIDER = re.compile(r"^;[-=\s]*$")

INPUT_HEADER = "작성할 대사 ← 여기에 입력"

CATEGORY = [
    ("KOJO_MESSAGE_COM", "조교 커맨드 대사", "성적"),
    ("DOG_KOJO", "수간(견화) 조교 대사", "성적"),
    ("COLOSSEUM_KOJO", "투기장 조교 대사", "성적/일반"),
    ("KOJO_MESSAGE_PALAMCNG", "파라미터 변화 반응", "혼합"),
    ("KOJO_MESSAGE_MARKCNG", "각인 변화 반응", "혼합"),
    ("SELF_KOJO", "자위 조교 대사", "성적"),
    ("DUNGEON_RYOUZYOKU", "던전 능욕", "성적"),
    ("BENKI_KOUJO", "변기 조교", "성적/스캇"),
    ("DUNGEON_VICTORY", "던전 전투 승리", "일반"),
    ("DUNGEON_ATTACK", "던전 전투", "일반"),
    ("NTR_KOUJO", "NTR", "성적"),
    ("EXUCUTION_KOUJO", "처형", "일반/고어"),
    ("MUSEUM_KOUJO", "박물관(전시)", "일반"),
    ("BANISHMENT_KOUJO", "추방", "일반"),
    ("PUBLIC_EXUCUTION", "공개 처형", "일반/고어"),
    ("GROTESQUE_KOUJO", "그로테스크/고어", "고어"),
    ("GOHOUBI", "포상", "성적/일반"),
    ("OSIOKI_KOUJO", "징벌", "성적/일반"),
    ("ENTERENEMY", "적대화", "일반"),
    ("GOBI_KOUJO", "어미(말버릇)", "일반"),
]

# 캐릭터 소개: Kn -> (표시명, 성격/컨셉, 성향, 말투/1인칭, 비고)
CHAR_INFO = {
    "K0": ("자애 (慈愛)",
           "자비롭고 상냥한 성모(聖母)형. 어떤 굴욕·역경에도 마왕을 원망하지 않고 용서·포용하며 희망을 잃지 않는다.",
           "받아들이는 헌신형. 저항보다 감싸안음.",
           "따뜻한 존댓말 '~어요/~네요', 상대를 걱정하는 어조.",
           ""),
    "K1": ("자신가 (自信家)",
           "오만하고 자신만만. 지기 싫어하고 도발적. 굴복당해도 허세와 자존심을 버리지 않는다.",
           "강한 척·도발. 굴복은 분한 듯.",
           "거만한 반말, 욕설('젠장') 섞임. 자신을 3인칭(SELF_CALL)으로 부르기도.",
           ""),
    "K2": ("소심 (気弱)",
           "겁많고 내성적. 늘 두려워하며 눈물·떨림이 많다.",
           "수동적·공포. 저항 못하고 애원.",
           "더듬는 말투, 말줄임표 다수. '싫어…', '무서워…', '용서해…'. 1인칭 SELF_CALL.",
           ""),
    "K3": ("고귀 (高貴)",
           "기품 있는 귀족 숙녀. 우아하고 자존심이 높다.",
           "품위 유지하려 애씀. 굴욕에 수치.",
           "정중한 여성어 '~예요/~어요', 우아한 어휘.",
           ""),
    "K4": ("냉철 (冷徹)",
           "냉정·침착·태연. 감정을 겉으로 드러내지 않고 담담하다.",
           "감정 억제·분석적. 동요를 감춤.",
           "담담한 평서체. 짧고 건조한 문장.",
           ""),
    "K5": ("마오 (マオ)",
           "유아적이고 순진한 어린 소녀. 언니(리리)를 찾고 있다.",
           "천진난만. 상황을 잘 이해 못함.",
           "어린애 말투, 짧은 문장. '언니…'를 자주 찾음.",
           "K11 리리와 자매 관계."),
    "K6": ("악녀 (悪女)",
           "요염하고 교활한 악녀. 독설과 저주를 내뱉는다.",
           "도발·조롱. 굴복해도 앙심.",
           "비웃는 반말, 저주·독설. 자신을 3인칭(SELF_CALL)으로.",
           ""),
    "K7": ("골드하트 (ゴールドハート)",
           "트럼프 사천왕 '하트'. 화려하고 예의 바른 무대 배우풍.",
           "화려·연극적. 우아한 굴복.",
           "정중한 '~습니다' 무대풍 어조.",
           "트럼프 4천왕(하트/스페이드/다이아/클럽) 중 하나."),
    "K8": ("실버스페이드 (シルバースペード)",
           "트럼프 사천왕 '스페이드'. 자존심과 복수심이 강한 기사·전사풍.",
           "저항적·투쟁적. 굴복을 굴욕으로.",
           "딱딱하고 강한 어조. 결의·분노.",
           "트럼프 4천왕 중 하나."),
    "K9": ("블랙다이아 (ブラックダイヤ)",
           "트럼프 사천왕 '다이아'. 강인한 전사형.",
           "호전적·강건. 정면돌파.",
           "굳세고 남성적인 어조.",
           "트럼프 4천왕 중 하나."),
    "K10": ("화이트클럽 (ホワイトクラブ)",
            "트럼프 사천왕 '클럽'. 조용하고 신비로운 마법사풍.",
            "차분·관조. 조용한 반응.",
            "조용조용한 어조, '우후 후' 웃음.",
            "트럼프 4천왕 중 하나."),
    "K11": ("리리 (リリィ)",
            "마을 처녀. 여동생 마오를 찾아 마왕성에 온 헌신적이고 강인한 언니.",
            "여동생 걱정이 최우선. 자신보다 마오.",
            "따뜻하지만 심지 굳은 어조. '마오만은…'.",
            "K5 마오와 자매 관계."),
    "K12": ("지적 (知的)",
            "지적이고 무뚝뚝한 학자·박사 타입. 모든 것을 관찰·연구 대상으로 본다.",
            "냉철한 관찰자. 자기 몸의 변화조차 분석.",
            "분석조 평서체 '~인데/~라니/관측된다', 전문용어.",
            ""),
    "K13": ("비호자 (庇護者)",
            "모성적이고 포용력 있는 누님. 역경 속에서도 마왕을 걱정하고 감싼다.",
            "감싸주는 모성형. 저항보다 포용.",
            "부드러운 누님 어조 '우후후♪', 2인칭 '당신'.",
            ""),
    "K14": ("귀공자 (貴公子)",
            "남성(소년~청년) 캐릭터. 육상 좋고 예의 바르나 긍지 높은 오만한 귀공자. (파일 헤더에 명시된 컨셉)",
            "긍지·자존심. 굴복을 수치로 여김.",
            "평소 오만한 귀족 어투; 애정(TALENT:85) 시 집사풍 존댓말 '~옵니다'.",
            "파일 헤더에 유일하게 컨셉 메모 존재: オトコ口上 / 育ちがよく礼儀正しい."),
}

CHAR_RE = re.compile(r"EVENT_(K\d+)_(.+?)\.ERB$", re.IGNORECASE)
INVALID_SHEET = re.compile(r"[:\\/?*\[\]]")


def categorize(func):
    f = func.lstrip("@")
    for key, label, owner in CATEGORY:
        if f.startswith(key):
            return label, owner
    for key, label, owner in CATEGORY:
        if key in f:
            return label, owner
    return "기타", "확인필요"


def char_key(fname):
    m = CHAR_RE.search(fname)
    return m.group(1) if m else ""


def char_label(fname):
    m = CHAR_RE.search(fname)
    if not m:
        return ""
    return f"{m.group(1)} {m.group(2)}"


def clean_comment(s):
    return s.lstrip(";").strip()


def scan_file(path):
    with open(path, "rb") as f:
        raw = f.read()
    lines = raw.decode("utf-8-sig").split("\n")

    fname = os.path.basename(path)
    character = char_label(fname)

    rows = []
    ordinal = 0
    func = ""
    section = ""
    last_comment = ""
    prev_kind = ""
    prev_comment_text = ""
    cond_stack = []
    last_written = ""

    for i, raw_line in enumerate(lines):
        line = raw_line.rstrip("\r")
        stripped = line.strip()

        if not stripped:
            prev_kind = "blank"
            continue

        if stripped.startswith(";"):
            if DIVIDER.match(stripped):
                if prev_kind == "comment" and prev_comment_text:
                    section = prev_comment_text
                prev_kind = "divider"
            else:
                c = clean_comment(stripped)
                if c:
                    last_comment = c
                    prev_comment_text = c
                prev_kind = "comment"
            continue

        if stripped.startswith("@"):
            func = stripped.split(",")[0].strip()
            section = ""
            cond_stack = []
            last_comment = ""
            prev_comment_text = ""
            prev_kind = "code"
            last_written = ""
            continue

        upper = stripped.upper()
        if upper.startswith("IF "):
            cond_stack.append((stripped[3:].strip(), last_comment))
        elif upper.startswith("ELSEIF "):
            if cond_stack:
                cond_stack.pop()
            cond_stack.append((stripped[7:].strip(), last_comment))
        elif upper == "ELSE":
            if cond_stack:
                cond_stack.pop()
            cond_stack.append(("(그 외의 경우)", last_comment))
        elif upper.startswith("ENDIF"):
            if cond_stack:
                cond_stack.pop()

        m = EMPTY_FORM.match(stripped)
        if m:
            keyword = m.group(1)
            branch_labels = [c for (_, c) in cond_stack if c]
            own = last_comment if last_comment and last_comment not in branch_labels else ""
            chain = " > ".join(branch_labels + ([own] if own else []))
            cond_top = cond_stack[-1][0] if cond_stack else ""
            rows.append({
                "id": f"{fname}#{ordinal:04d}",
                "file": fname,
                "character": character,
                "func": func,
                "section": section,
                "branch": chain,
                "cond": cond_top,
                "prev": last_written,
                "line": i + 1,
                "ordinal": ordinal,
                "keyword": keyword,
            })
            ordinal += 1
            last_comment = ""
            prev_kind = "code"
            continue

        if upper.startswith("PRINTFORM") or upper.startswith("PRINTSINGLEFORM"):
            payload = stripped.split(None, 1)
            if len(payload) == 2 and payload[1].strip():
                last_written = payload[1].strip()

        last_comment = ""
        prev_kind = "code"

    return rows


# ---- 스타일 ----
HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(bold=True, color="FFFFFF")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
TITLE_FILL = PatternFill("solid", fgColor="1F3864")
TITLE_FONT = Font(bold=True, color="FFFFFF", size=13)
INTRO_KEY_FILL = PatternFill("solid", fgColor="D6DCE4")
INTRO_KEY_FONT = Font(bold=True, color="1F3864")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical="top", wrap_text=True)
WRAP_MID = Alignment(vertical="center", wrap_text=True)

HEADERS = [
    ("ID", 20), ("파일", 26), ("캐릭터", 14), ("카테고리", 18),
    ("담당힌트", 10), ("함수", 26), ("커맨드/장면", 26), ("상황(분기)", 34),
    ("조건식(원문)", 34), ("직전 작성 예시(톤 참고)", 40),
    (INPUT_HEADER, 60),
    ("_line", 8), ("_ordinal", 9), ("_keyword", 16),
]
NCOL = len(HEADERS)
INPUT_COL = 11


def intro_lines(fname):
    key = char_key(fname)
    if key in CHAR_INFO:
        label, concept, tend, tone, note = CHAR_INFO[key]
        return label, [
            ("성격/컨셉", concept),
            ("성향", tend),
            ("말투/1인칭", tone),
        ] + ([("비고", note)] if note else [])
    # 캐릭터 전용이 아닌 파일
    return fname, [
        ("성격/컨셉", "특정 캐릭터 전용이 아닌 시스템/공용 이벤트 파일. 각 행의 '카테고리'·'상황(분기)'을 보고 맥락에 맞춰 작성."),
    ]


def write_sheet(wb, title, fname, rows):
    ws = wb.create_sheet(title)
    label, intro = intro_lines(fname)

    # 제목
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=NCOL)
    t = ws.cell(1, 1, f"🎭 {label} — 대사 작성 시트   (빈 자리 {len(rows)}개)")
    t.fill = TITLE_FILL
    t.font = TITLE_FONT
    t.alignment = WRAP_MID
    ws.row_dimensions[1].height = 24

    r = 2
    for k, v in intro:
        kc = ws.cell(r, 1, k)
        kc.fill = INTRO_KEY_FILL
        kc.font = INTRO_KEY_FONT
        kc.alignment = WRAP_MID
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=NCOL)
        vc = ws.cell(r, 2, v)
        vc.alignment = WRAP_MID
        r += 1

    # 안내 한 줄
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=NCOL)
    g = ws.cell(r, 1, "노란 칸(작성할 대사)에만 입력하세요. 한 행 = 대사 1줄. "
                      "_line/_ordinal/_keyword(숨김)은 주입 좌표라 수정·삭제 금지. "
                      "빈 행은 자동으로 건너뜁니다.")
    g.font = Font(italic=True, color="C00000")
    g.alignment = WRAP_MID
    r += 1

    # 빈 줄
    r += 1

    # 헤더
    header_row = r
    for c, (name, width) in enumerate(HEADERS, start=1):
        cell = ws.cell(header_row, c, name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = width
    r += 1

    # 데이터
    for row in rows:
        cat, owner = categorize(row["func"])
        vals = [
            row["id"], row["file"], row["character"], cat, owner, row["func"],
            row["section"], row["branch"], row["cond"], row["prev"],
            "", row["line"], row["ordinal"], row["keyword"],
        ]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(r, c, v)
            cell.alignment = WRAP
            cell.border = BORDER
            if c == INPUT_COL:
                cell.fill = INPUT_FILL
        r += 1

    ws.freeze_panes = ws.cell(header_row + 1, 1).coordinate
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(NCOL)}{r - 1}"
    for col in (12, 13, 14):
        ws.column_dimensions[get_column_letter(col)].hidden = True
    return len(rows)


def sheet_title(fname, used):
    lab = char_label(fname)
    base = lab if lab else os.path.splitext(fname)[0]
    base = INVALID_SHEET.sub("_", base)[:31].strip() or "sheet"
    title = base
    n = 2
    while title in used:
        suffix = f"~{n}"
        title = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(title)
    return title


def main():
    if not os.path.isdir(ERB_DIR):
        print(f"ERB 폴더 없음: {ERB_DIR}", file=sys.stderr)
        sys.exit(1)

    files = sorted(f for f in os.listdir(ERB_DIR) if f.upper().endswith(".ERB"))
    per_file = {}
    scanned = {}
    for f in files:
        rows = scan_file(os.path.join(ERB_DIR, f))
        if rows:
            per_file[f] = len(rows)
            scanned[f] = rows

    wb = Workbook()
    wb.remove(wb.active)

    # 안내 시트
    guide = wb.create_sheet("안내")
    guide.column_dimensions["A"].width = 100
    guide_lines = [
        "■ 대사 작성 안내",
        "",
        "· 이 파일은 '캐릭터별 시트'로 되어 있습니다. 아래 탭(시트)에서 채우려는 캐릭터를 고르세요.",
        "· 각 시트 맨 위에 그 캐릭터의 성격/성향/말투 소개가 있습니다. 그 톤에 맞춰 대사를 쓰면 됩니다.",
        "· 노란색 '작성할 대사 ← 여기에 입력' 칸에만 대사를 씁니다. 다른 칸은 참고용입니다.",
        "· 한 행 = 대사 한 줄(PRINTFORM 1개). 셀 안에서 줄바꿈하면 주입 시 공백으로 합쳐집니다.",
        "· 비워 둔 행은 자동으로 건너뜁니다. 필요한 것만 채워도 됩니다.",
        "· 숨겨진 _line / _ordinal / _keyword 열은 '원본의 정확한 자리'를 가리키는 좌표입니다. 수정·삭제하지 마세요.",
        "· 시트/행을 지우거나 정렬해도 좌표(_line·_keyword)만 살아 있으면 정확히 주입됩니다.",
        "",
        "■ 다 채운 뒤 주입(되돌리기)",
        "    python tools/fill-list/apply_fill_list.py --dry-run   (미리보기)",
        "    python tools/fill-list/apply_fill_list.py             (실제 주입)",
        "",
        "· 안전장치: 대상 라인이 여전히 '빈 자리'이고 명령어가 일치할 때만 덮어씁니다. 어긋나면 건너뛰고 보고합니다.",
        "· 원본의 UTF-8 BOM / CRLF 줄끝 / 탭 들여쓰기 / 라인 수를 그대로 보존합니다.",
    ]
    for i, ln in enumerate(guide_lines, start=1):
        cell = guide.cell(i, 1, ln)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if ln.startswith("■"):
            cell.font = Font(bold=True, size=12, color="1F3864")

    # 요약 시트
    summary = wb.create_sheet("요약")
    summary.append(["시트(캐릭터/파일)", "파일", "빈 자리 수"])
    for c in summary[1]:
        c.font = Font(bold=True)
    summary.column_dimensions["A"].width = 26
    summary.column_dimensions["B"].width = 30
    summary.column_dimensions["C"].width = 12

    used_titles = set()
    total = 0
    # 슬롯 많은 순으로 시트 생성
    for f in sorted(per_file, key=lambda k: -per_file[k]):
        title = sheet_title(f, used_titles)
        cnt = write_sheet(wb, title, f, scanned[f])
        summary.append([title, f, cnt])
        total += cnt
    summary.append(["합계", "", total])
    summary[summary.max_row][0].font = Font(bold=True)

    wb.save(OUT_XLSX)
    print(f"OK: {total} slots / {len(per_file)} sheets -> {OUT_XLSX}")


if __name__ == "__main__":
    main()
